"""Split one provider sheet of the customer's consumption workbook into raw monthly CSVs.

The workbook stores usage_start_date/usage_end_date as real dates and cost as a real
number - "2026-6" and "$12,499" are just Excel's *display* formatting (number_format),
not the underlying values. clean_date()/clean_cost() normalize both that clean case and
a messier text export (e.g. a hand-edited copy where a cell literally contains "2026-6"
or "$12,499") so the split step is robust either way.
"""
from __future__ import annotations

import csv
import logging
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

from .config import COST_COLUMN, DATE_COLUMNS, MONTH_ABBR, MONTH_KEY_COLUMN

logger = logging.getLogger(__name__)


def clean_date(value) -> date | None:
    """Normalize a workbook date cell to a date. Accepts real datetimes/dates, or
    text fallbacks like '2026-6', '2026-06-01', '01/06/26', '01/06/2026'."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    if text == "":
        return None
    if re.match(r"^\d{1,2}/\d{1,2}/\d{2}$", text):
        return datetime.strptime(text, "%d/%m/%y").date()
    if re.match(r"^\d{1,2}/\d{1,2}/\d{4}$", text):
        return datetime.strptime(text, "%d/%m/%Y").date()
    if re.match(r"^\d{4}-\d{1,2}-\d{1,2}$", text):
        return datetime.strptime(text, "%Y-%m-%d").date()
    if re.match(r"^\d{4}-\d{1,2}$", text):
        year_str, month_str = text.split("-")
        return date(int(year_str), int(month_str), 1)

    raise ValueError(f"Unrecognized date value: {value!r}")


def clean_cost(value) -> float:
    """Normalize a workbook cost cell to a float. Accepts real numbers, or text
    fallbacks like '$12,499' / '12,499.00'."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace("$", "").replace(",", "")
    return float(text) if text else 0.0


def read_provider_sheet(workbook_path: Path, provider: str):
    """Read one provider sheet, clean its date/cost columns, and group rows by
    (year, month) of usage_start_date.

    Returns (headers, date_col_indices, cost_col_idx, months) where:
      headers: list[str] mirroring the sheet's used column range (blanks kept as "")
      date_col_indices: {column_name: index} for whichever of DATE_COLUMNS are present
      cost_col_idx: index of the cost column, or None if absent
      months: {(year, month): [row_values, ...]} with dates/cost already cleaned
    """
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if provider not in workbook.sheetnames:
        raise ValueError(f"Sheet '{provider}' not found in workbook {workbook_path}")
    worksheet = workbook[provider]

    rows_iter = worksheet.iter_rows(values_only=True)
    headers = ["" if cell is None else str(cell) for cell in next(rows_iter)]

    if MONTH_KEY_COLUMN not in headers:
        raise ValueError(f"Sheet '{provider}' is missing the '{MONTH_KEY_COLUMN}' column")

    date_col_indices = {name: headers.index(name) for name in DATE_COLUMNS if name in headers}
    cost_col_idx = headers.index(COST_COLUMN) if COST_COLUMN in headers else None
    month_key_idx = date_col_indices[MONTH_KEY_COLUMN]

    months: dict[tuple[int, int], list[list]] = {}
    skipped = 0
    for row in rows_iter:
        if row is None or all(v is None for v in row):
            continue
        row = list(row)

        for idx in date_col_indices.values():
            if idx < len(row):
                try:
                    row[idx] = clean_date(row[idx])
                except ValueError:
                    row[idx] = None

        usage_date = row[month_key_idx] if month_key_idx < len(row) else None
        if usage_date is None:
            skipped += 1
            continue

        if cost_col_idx is not None and cost_col_idx < len(row):
            row[cost_col_idx] = clean_cost(row[cost_col_idx])

        months.setdefault((usage_date.year, usage_date.month), []).append(row)

    if skipped:
        logger.warning(
            "Sheet '%s': skipped %d row(s) with no usable %s", provider, skipped, MONTH_KEY_COLUMN
        )

    return headers, date_col_indices, cost_col_idx, months


def write_month_csv(
    headers: list[str],
    date_col_indices: dict[str, int],
    cost_col_idx: int | None,
    rows: list[list],
    provider: str,
    year: int,
    month: int,
    output_dir: Path,
) -> Path:
    """Write one month's rows to <output_dir>/<provider>/<year>/<month:02d>/<provider><Mon>.csv,
    formatting dates as DD/MM/YY and cost as a plain number (e.g. 12499.00)."""
    month_dir = output_dir / provider / str(year) / f"{month:02d}"
    month_dir.mkdir(parents=True, exist_ok=True)
    out_path = month_dir / f"{provider}{MONTH_ABBR[month]}.csv"

    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for row in rows:
            formatted = list(row)

            for idx in date_col_indices.values():
                if idx < len(formatted) and formatted[idx] is not None:
                    formatted[idx] = formatted[idx].strftime("%d/%m/%y")

            if cost_col_idx is not None and cost_col_idx < len(formatted) and formatted[cost_col_idx] is not None:
                formatted[cost_col_idx] = f"{formatted[cost_col_idx]:.2f}"

            if len(formatted) < len(headers):
                formatted += [""] * (len(headers) - len(formatted))
            else:
                formatted = formatted[: len(headers)]

            writer.writerow(["" if v is None else v for v in formatted])

    return out_path
